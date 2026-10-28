"""Rebuild Byron's FY2028 Budget Instructions content on the Maziv Group A4 master.

Content source : Byron's deck (every paragraph, run, table cell and note, verbatim).
Look source    : the Group master + curated layouts carried in the same file,
                 styled per maziv-brand-palette (light palette, #7C1C91 on white).
Usage: python3 -I build_deck.py SRC.pptx OUT.pptx
"""
import sys
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]

# ---- Maziv brand palette (canonical, Sep 2026) ----------------------------
PURPLE, PURPLE_L, TINT, TINT2, BADGE = "7C1C91", "984DA9", "F4ECF7", "FAF6FC", "E8D5F0"
INK, INK_S, MUTE, LINE, PAPER = "3D4552", "1F2530", "8B9097", "E3E3E3", "FFFFFF"
AMBER_F, AMBER_T = "FDF1DC", "8F5F07"

FOOTER_TEXT = "Maziv Group  •  FY2028 Budget Instructions  •  Strictly confidential"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

src = Presentation(SRC)          # read-only content source
prs = Presentation(SRC)          # host: keeps master, layouts, theme, notes master
S = src.slides
N_SLIDES = len(S)

# ---- strip Byron's slides from the host -----------------------------------
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

master = prs.slide_masters[0]
LAYOUT = {l.name: l for l in master.slide_layouts}
L_COVER = LAYOUT["Cover 1 - Light"]
L_QUAD = LAYOUT["7_Title and Content"]
L_PANEL = LAYOUT["10_Title and Content"]

report = []  # formatting decisions worth telling the user


# ---- content helpers ---------------------------------------------------------
def runs_of(paragraph):
    """Merge adjacent runs with the same bold state; keep text verbatim."""
    out = []
    for r in paragraph.runs:
        b = bool(r.font.bold)
        hl = r._r.find(f".//{{{A_NS}}}highlight") is not None
        if out and out[-1][1] == b and out[-1][2] == hl:
            out[-1][0] += r.text
        else:
            out.append([r.text, b, hl])
    return [(t, b, h) for t, b, h in out]


def paras_of(text_frame):
    return [(p.level, runs_of(p)) for p in text_frame.paragraphs if "".join(r.text for r in p.runs)]


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def set_bullet(p, mode, marL=171450, indent=-171450):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if mode == "none":
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))
    elif mode == "char":
        pPr.set("marL", str(marL))
        pPr.set("indent", str(indent))
        bf = etree.SubElement(pPr, qn("a:buFont"))
        bf.set("typeface", "Arial")
        bc = etree.SubElement(pPr, qn("a:buChar"))
        bc.set("char", "•")


def fill_tf(tf, paras, size, color, bold=None, bullet=None, space_after=None,
            align=None, heading=None):
    """heading = (size, color) applied to the first paragraph, bold, no bullet."""
    tf.clear()
    first = True
    for i, (lvl, runs) in enumerate(paras):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        if align is not None:
            p.alignment = align
        if space_after is not None:
            p.space_after = Pt(space_after)
        is_head = heading is not None and i == 0
        for text, b, _hl in runs:
            r = p.add_run()
            r.text = text
            if is_head:
                r.font.size, r.font.bold = Pt(heading[0]), True
                r.font.color.rgb = RGBColor.from_string(heading[1])
            else:
                r.font.size = Pt(size)
                r.font.color.rgb = RGBColor.from_string(color)
                r.font.bold = bold if bold is not None else (True if b else None)
        if is_head:
            set_bullet(p, "none")
        elif bullet:
            set_bullet(p, bullet)
    return tf


def textbox(slide, name, x, y, w, h, margin=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    return tb


def solid(shape, hexcol, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hexcol)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False


def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(idx)


def set_title(slide, src_slide):
    t = ph(slide, 0)
    t.text_frame.clear()
    r = t.text_frame.paragraphs[0].add_run()
    r.text = "".join(rr.text for p in src_slide.shapes.title.text_frame.paragraphs for rr in p.runs)
    t.name = src_slide.shapes.title.name


# ---- table helpers -----------------------------------------------------------
def _ln(side, color=None, w_pt=0.75):
    ln = etree.Element(qn(f"a:ln{side}"))
    if color is None:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    else:
        ln.set("w", str(int(w_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
    return ln


def cell_borders(cell, bottom=LINE):
    tcPr = cell._tc.get_or_add_tcPr()
    for side in ("L", "R", "T", "B"):
        for old in tcPr.findall(qn(f"a:ln{side}")):
            tcPr.remove(old)
    for i, side in enumerate(("L", "R", "T", "B")):
        tcPr.insert(i, _ln(side, bottom if side == "B" else None))


def cell_fill(cell, hexcol):
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(hexcol)


def add_table(slide, src_table_shape, x, y, col_widths, head_h, row_h, head_size, body_size,
              first_col_bold=True):
    st = src_table_shape.table
    nrows, ncols = len(st.rows), len(st.columns)
    total_h = head_h + row_h * (nrows - 1)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(col_widths)), Inches(total_h))
    gf.name = src_table_shape.name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for ci, cw in enumerate(col_widths):
        tbl.columns[ci].width = Inches(cw)
    for ri in range(nrows):
        tbl.rows[ri].height = Inches(head_h if ri == 0 else row_h)
        for ci in range(ncols):
            scell = st.cell(ri, ci)
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            paras = [(p.level, runs_of(p)) for p in scell.text_frame.paragraphs]
            if not paras:
                paras = [(0, [])]
            highlighted = any(h for _, rs in paras for _, _, h in rs)
            if ri == 0:
                fill_hex, col, bold_all, size = PURPLE, PAPER, True, head_size
            else:
                fill_hex = PAPER if ri % 2 else TINT2
                col = INK_S if (ci == 0 and first_col_bold) else INK
                bold_all = True if (ci == 0 and first_col_bold) else None
                size = body_size
            if highlighted:
                fill_hex, col, bold_all = AMBER_F, AMBER_T, True
                report.append(f"Slide {slide.slide_id}: Byron's yellow highlight on '{scell.text}' "
                              f"rendered as the brand amber 'watch' cell ({AMBER_F}/{AMBER_T}).")
            cell_fill(cell, fill_hex)
            cell_borders(cell, bottom=LINE if ri else PURPLE)
            first = True
            for lvl, runs in paras:
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                for text, b, _hl in runs:
                    r = p.add_run()
                    r.text = text
                    if ri and text.strip() == "NEW":          # Byron's NEW tag → purple badge text
                        r.font.size, r.font.bold = Pt(max(body_size - 1.5, 8)), True
                        r.font.color.rgb = RGBColor.from_string(PURPLE)
                    else:
                        r.font.size = Pt(size)
                        r.font.color.rgb = RGBColor.from_string(col)
                        r.font.bold = bold_all if bold_all is not None else (True if b else None)
    return gf


# ---- layout chrome: footer line + page number on the two content layouts --------
def add_layout_chrome(layout, sid_base):
    tmp = prs.slides.add_slide(layout)
    # footer line
    tb = textbox(tmp, "Footer line", 1.5, 7.66, 6.6, 0.2)
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = FOOTER_TEXT
    r.font.size, r.font.italic = Pt(8), True
    r.font.color.rgb = RGBColor.from_string(MUTE)
    # page number  n / N
    nb = textbox(tmp, "Page number", 9.48, 7.66, 1.8, 0.2)
    p2 = nb.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p2._p, qn("a:fld"))
    fld.set("id", "{B6F15528-21DB-4D62-9E3B-7A1C0A5D2F10}")
    fld.set("type", "slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"))
    rPr.set("lang", "en-ZA")
    rPr.set("sz", "800")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", MUTE)
    etree.SubElement(fld, qn("a:t")).text = "1"
    r2 = p2.add_run()
    r2.text = f" / {N_SLIDES}"
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGBColor.from_string(MUTE)
    # move both shapes onto the layout, then drop the temp slide
    spTree = layout.shapes._spTree
    for i, shp in enumerate((tb, nb)):
        el = shp._element
        el.getparent().remove(el)
        el.nvSpPr.cNvPr.id = sid_base + i
        spTree.append(el)
    sid = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(sid.rId)
    prs.slides._sldIdLst.remove(sid)


add_layout_chrome(L_QUAD, 9001)
add_layout_chrome(L_PANEL, 9001)

# ============================================================================
# SLIDE 1 — cover (Cover 1 - Light)
# ============================================================================
s1 = prs.slides.add_slide(L_COVER)
src1 = S[0]
cover_paras = paras_of(shape_by_name(src1, "Subtitle 2").text_frame)   # two paragraphs
t = ph(s1, 10)
t.name = "Subtitle 2"
fill_tf(t.text_frame, cover_paras[:1], 40, PURPLE, bold=True, bullet="none")
sub = ph(s1, 11)
sub.name = "Subtitle 2 (line 2)"
fill_tf(sub.text_frame, cover_paras[1:], 24, PURPLE, bullet="none")

# ============================================================================
# SLIDE 2 — four principles (7_Title and Content)
# ============================================================================
s2 = prs.slides.add_slide(L_QUAD)
src2 = S[1]
set_title(s2, src2)
quad = [("Text Placeholder 2", 1, "Content Placeholder 3", 2),
        ("Text Placeholder 4", 15, "Content Placeholder 5", 16),
        ("Text Placeholder 6", 17, "Content Placeholder 7", 18),
        ("Text Placeholder 8", 19, "Content Placeholder 9", 20)]
for hname, hidx, cname, cidx in quad:
    h = ph(s2, hidx)
    h.name = hname
    fill_tf(h.text_frame, paras_of(shape_by_name(src2, hname).text_frame), 14, PURPLE, bold=True, bullet="none")
    c = ph(s2, cidx)
    c.name = cname
    lp = next(p for p in L_QUAD.placeholders if p.placeholder_format.idx == cidx)
    c.left, c.width = lp.left, lp.width          # set all four, or the xfrm gets a zero width
    c.top, c.height = Inches(1.80 if lp.top < Inches(3) else 5.10), Inches(2.30)
    c.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    fill_tf(c.text_frame, paras_of(shape_by_name(src2, cname).text_frame), 11, INK, space_after=5)

# ============================================================================
# SLIDE 3 — process flow + rules panel (10_Title and Content)
# ============================================================================
s3 = prs.slides.add_slide(L_PANEL)
src3 = S[2]
set_title(s3, src3)
xs = [0.62, 2.75, 4.88, 7.01, 9.14]
for i, x in enumerate(xs, start=1):
    arrow_src = shape_by_name(src3, f"Step {i} arrow")
    shp = s3.shapes.add_shape(MSO_SHAPE.PENTAGON if i == 1 else MSO_SHAPE.CHEVRON,
                              Inches(x), Inches(1.40), Inches(2.05), Inches(0.75))
    shp.name = arrow_src.name
    solid(shp, PURPLE if i % 2 else PURPLE_L)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.15)
    fill_tf(tf, paras_of(arrow_src.text_frame), 12, PAPER, bold=True, bullet="none", align=PP_ALIGN.CENTER)
    txt_src = shape_by_name(src3, f"Step {i} text")
    tb = textbox(s3, txt_src.name, x, 2.35, 1.93, 1.55)
    fill_tf(tb.text_frame, paras_of(txt_src.text_frame), 11, INK, bullet="none", space_after=3,
            heading=(12, PURPLE))

panel = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.62), Inches(4.15), Inches(10.43), Inches(2.30))
panel.name = "Rules panel"
panel.adjustments[0] = 0.03
solid(panel, TINT)
for name, x in (("Rules left", 0.85), ("Rules right", 6.0)):
    tb = textbox(s3, name, x, 4.38, 4.80, 2.2)
    fill_tf(tb.text_frame, paras_of(shape_by_name(src3, name).text_frame), 12, INK, bullet="char",
            space_after=6, heading=(14, PURPLE))
tag = textbox(s3, "Draft tag", 7.3, 7.15, 3.75, 0.2)
fill_tf(tag.text_frame, paras_of(shape_by_name(src3, "Draft tag").text_frame), 8, MUTE, bullet="none",
        align=PP_ALIGN.RIGHT)
s3.notes_slide.notes_text_frame.text = src3.notes_slide.notes_text_frame.text

# ============================================================================
# SLIDES 4–8 — tables (10_Title and Content)
# ============================================================================
TABLE_SPECS = {
    4: dict(x=0.62, y=1.28, col_widths=[4.60, 5.83], head_h=0.40, row_h=0.375, head_size=10.5, body_size=10.5),
    5: dict(x=0.62, y=1.28, col_widths=[3.70, 2.45, 4.28], head_h=0.40, row_h=0.46, head_size=10.5, body_size=10.5),
    6: dict(x=0.62, y=1.28, col_widths=[5.45, 2.20, 2.78], head_h=0.40, row_h=0.42, head_size=10.5, body_size=10.5),
    7: dict(x=0.62, y=1.28, col_widths=[5.43, 2.00, 3.00], head_h=0.40, row_h=0.62, head_size=10.5, body_size=10.5),
    8: dict(x=0.62, y=1.28, col_widths=[3.30, 7.13], head_h=0.40, row_h=0.54, head_size=10.5, body_size=10.5),
}
for n in range(4, 9):
    s = prs.slides.add_slide(L_PANEL)
    srcs = S[n - 1]
    set_title(s, srcs)
    add_table(s, shape_by_name(srcs, "Table"), **TABLE_SPECS[n])
    if n == 7:
        fn_src = shape_by_name(srcs, "Footnote")
        fn = textbox(s, fn_src.name, 0.62, 7.05, 10.43, 0.25)
        fill_tf(fn.text_frame, paras_of(fn_src.text_frame), 9, MUTE, bullet="none")
    if srcs.has_notes_slide and srcs.notes_slide.notes_text_frame.text.strip():
        s.notes_slide.notes_text_frame.text = srcs.notes_slide.notes_text_frame.text

prs.save(OUT)
print(f"wrote {OUT}: {len(prs.slides)} slides")
for line in report:
    print("NOTE:", line)
